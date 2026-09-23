"""Tests for fishbehavior.catalog (workbook cleaning + video matching).

Everything is synthetic: the workbook is written here with the same header text as the
real one (line breaks, the 'Compund' typo) but invented compounds and values, and the
"videos" are empty files — only their names matter for matching.
"""

import math
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from fishbehavior.catalog import (
    CatalogError,
    build_catalog,
    clean_trials,
    match_videos,
    merge_split_recordings,
    normalize_header,
    read_workbook,
    scan_videos,
)
from fishbehavior.cli import main
from fishbehavior.config import load_settings

# Header row exactly as it appears in the real workbook (text only, no data).
HEADERS = [
    "Date of EXP:", "Subject #:", "Strain:", "Sex (M/F):", "Age (~):", "Compund:", "Conc. (mM):",
    "Agent Exposure Time (min):", "NTT \nTime (min):", "UV \nExposure\nTime (min)",
    "TDM (Full Arena):", "TDM (Top Half):", "TDM (Bot Half):",
    "Velocity (Full Arena)", "Velocity (Top Half)", "Velocity (Bot Half)",
    "Time Spent (Top):", "Time Spent (Bot):",
    "H2O (Before):", "H2O (After):", "Brain Tissue:", "Body Tissue:",
]

TRACKED = [100.0, 60.0, 40.0, 2.0, 1.5, 2.5, 300.0, 300.0]  # the 8 movement values
TOP_ZERO = [100.0, "-", 100.0, 2.0, "-", 2.0, 0, 600.0]  # "-" = no time in the top half
ALL_NA = ["N/A"] * 8
ALL_BLANK = [None] * 8


def row(subject, sex, compound, movement, conc=0.03, uv=10, date=240109):
    """One workbook row with invented values (only what the tests care about varies)."""
    return [date, subject, "Strain-X", sex, 1.5, compound, conc, 20, 10, uv, *movement,
            "AVAIL", "AVAIL", "N/A", "N/A"]


def make_workbook(path: Path, rows) -> Path:
    """Write a workbook with the real header row and the given data rows."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    for values in rows:
        sheet.append(values)
    book.save(path)
    return path


def touch(folder: Path, *names: str) -> None:
    """Create empty stand-in "video" files."""
    folder.mkdir(parents=True, exist_ok=True)
    for name in names:
        (folder / name).write_bytes(b"")


PATTERN = load_settings(environ={}).params["catalog"]["video_name_pattern"]
EXTENSIONS = [".mp4", ".avi", ".mov", ".mkv"]


# --- headers -------------------------------------------------------------------


def test_normalize_header_removes_colons_and_line_breaks():
    assert normalize_header("NTT \nTime (min):") == "ntt time (min)"
    assert normalize_header("UV \nExposure\nTime (min)") == "uv exposure time (min)"


def test_read_workbook_maps_headers_and_keeps_raw_text(tmp_path):
    book = make_workbook(tmp_path / "db.xlsx", [row("0001", "M", "Drug-A", ALL_NA, uv="10*")])

    table, unknown = read_workbook(book)

    assert unknown == []
    assert {"subject_id", "compound", "tdm_full", "uv_min", "body_tissue"} <= set(table.columns)
    assert table.loc[0, "subject_id"] == "0001"  # leading zeros survive
    assert table.loc[0, "tdm_full"] == "N/A"  # not silently turned into "missing" by pandas
    assert table.loc[0, "uv_min"] == "10*"


def test_missing_required_column_is_a_clear_error(tmp_path):
    book = openpyxl.Workbook()
    book.active.append(["Subject #:", "Sex (M/F):"])  # no compound, no movement columns
    book.save(tmp_path / "bad.xlsx")

    with pytest.raises(CatalogError, match="required column"):
        read_workbook(tmp_path / "bad.xlsx")


def test_unknown_columns_are_kept_and_reported(tmp_path):
    book = openpyxl.Workbook()
    book.active.append([*HEADERS, "Extra Note:"])
    book.active.append([*row("0001", "M", "Drug-A", TRACKED), "hello"])
    book.save(tmp_path / "db.xlsx")

    table, unknown = read_workbook(tmp_path / "db.xlsx")

    assert unknown == ["Extra Note:"]
    assert table.loc[0, "extra_note"] == "hello"


# --- cleaning rules (scope §5) --------------------------------------------------------


def cleaned(tmp_path, rows):
    """Read + clean a synthetic workbook; returns (table indexed by subject_id, issues)."""
    table, _ = read_workbook(make_workbook(tmp_path / "db.xlsx", rows))
    result, issues = clean_trials(table)
    return result.set_index("subject_id"), issues


def test_blank_compound_rows_are_dropped(tmp_path):
    table, _ = cleaned(tmp_path, [row("0001", "M", "Drug-A", TRACKED), row("0002", "F", None, ALL_BLANK)])
    assert list(table.index) == ["0001"]


def test_body_tissue_is_dropped(tmp_path):
    table, _ = cleaned(tmp_path, [row("0001", "M", "Drug-A", TRACKED)])
    assert "body_tissue" not in table.columns


def test_dash_in_zone_columns_means_zero(tmp_path):
    table, issues = cleaned(tmp_path, [row("0001", "M", "Drug-A", TOP_ZERO)])
    assert table.loc["0001", "tdm_top"] == 0.0
    assert table.loc["0001", "velocity_top"] == 0.0
    assert table.loc["0001", "ntt_tracked"]
    assert issues["Unexpected text in movement columns"] == []


@pytest.mark.parametrize("movement", [ALL_NA, ALL_BLANK], ids=["N/A", "blank"])
def test_all_eight_missing_means_untracked_and_stays_missing(tmp_path, movement):
    table, _ = cleaned(tmp_path, [row("0001", "M", "Drug-A", movement)])
    assert not table.loc["0001", "ntt_tracked"]
    assert table.loc["0001", ["tdm_full", "time_bottom_s"]].isna().all()  # masked, never imputed


def test_partly_missing_movement_is_still_tracked(tmp_path):
    table, _ = cleaned(tmp_path, [row("0001", "M", "Drug-A", [*TRACKED[:7], "N/A"])])
    assert table.loc["0001", "ntt_tracked"]
    assert math.isnan(table.loc["0001", "time_bottom_s"])


def test_unexpected_text_is_missing_and_reported(tmp_path):
    table, issues = cleaned(tmp_path, [row("0001", "M", "Drug-A", ["oops", *TRACKED[1:]])])
    assert math.isnan(table.loc["0001", "tdm_full"])
    assert issues["Unexpected text in movement columns"] == ["Excel row 2: tdm_full = 'oops' (treated as missing)"]


def test_uv_time_with_asterisk_becomes_missing_but_raw_text_is_kept(tmp_path):
    table, _ = cleaned(tmp_path, [row("0001", "M", "Drug-A", TRACKED, uv="10*")])
    assert math.isnan(table.loc["0001", "uv_min"])
    assert table.loc["0001", "uv_min_raw"] == "10*"


def test_dates_subjects_sex_and_concentration_are_normalized(tmp_path):
    table, _ = cleaned(tmp_path, [
        row(7, "m ", "Drug-A", TRACKED, conc="0.03 + 0.01", date=240109),  # subject stored as a number
        row("0008", "F", "Drug-B", TRACKED, conc=0.1, date=251231),
    ])
    assert list(table.index) == ["0007", "0008"]  # zero-padded like the video names
    assert table.loc["0007", "subject_num"] == 7
    assert table.loc["0007", "sex"] == "M"
    assert table.loc["0007", "exp_date"] == "2024-01-09"
    assert table.loc["0007", "concentration_raw"] == "0.03 + 0.01"
    assert math.isnan(table.loc["0007", "concentration_mm"])  # combination dose: text only
    assert table.loc["0008", "concentration_mm"] == 0.1


def test_bad_subject_and_sex_values_are_reported(tmp_path):
    table, issues = cleaned(tmp_path, [row("abc", "M", "Drug-A", TRACKED), row("0002", "X", "Drug-A", TRACKED)])
    assert list(table.index) == ["0002"]  # unreadable subject row skipped
    assert issues["Unreadable subject numbers"] == ["Excel row 2: subject 'abc' is not a whole number (row skipped)"]
    assert issues["Unexpected sex values"] == ["Excel row 3: sex 'X' is not M or F"]


# --- split recordings ---------------------------------------------------------------


def test_repeated_subject_is_joined_into_one_split_recording(tmp_path):
    table, _ = read_workbook(make_workbook(tmp_path / "db.xlsx", [
        row("0005", "M", "Drug-A", ALL_BLANK), row("0005", "M", "Drug-A", ALL_BLANK), row("0006", "F", "Drug-A", TRACKED),
    ]))
    merged, split_notes, conflicts = merge_split_recordings(clean_trials(table)[0])

    assert list(merged["subject_id"]) == ["0005", "0006"]
    first = merged.iloc[0]
    assert first["n_workbook_rows"] == 2 and first["excel_rows"] == "2;3"
    assert len(split_notes) == 1 and conflicts == []
    assert merged["ntt_tracked"].dtype == bool  # column types survive the join


def test_conflicting_values_between_parts_are_reported(tmp_path):
    table, _ = read_workbook(make_workbook(tmp_path / "db.xlsx", [
        row("0005", "M", "Drug-A", ALL_BLANK, conc=0.03), row("0005", "M", "Drug-A", ALL_BLANK, conc=0.1),
    ]))
    merged, _, conflicts = merge_split_recordings(clean_trials(table)[0])

    assert merged.loc[0, "concentration_mm"] == 0.03  # first row wins
    assert any("concentration_raw differs" in c for c in conflicts)


# --- video names -----------------------------------------------------------------------


def test_scan_videos_parses_names_parts_and_skips_hidden_files(tmp_path):
    touch(tmp_path, "F_0042.mp4", "copy-M_0001.MP4", "M_0012a.mp4", "M_0012_b.avi", "M_0013-A.mov",
          "._F_0042.mp4", "notes.txt", "random.mp4")
    touch(tmp_path / "sub", "F_0002.mkv")  # sub-folders are searched too

    videos = scan_videos(tmp_path, PATTERN, EXTENSIONS).set_index("file_name")

    assert set(videos.index) == {"F_0042.mp4", "copy-M_0001.MP4", "M_0012a.mp4", "M_0012_b.avi",
                                 "M_0013-A.mov", "random.mp4", "F_0002.mkv"}
    assert videos.loc["F_0042.mp4", ["subject_num", "part"]].tolist() == [42, ""]
    assert videos.loc["copy-M_0001.MP4", "subject_num"] == 1  # prefix tolerated
    assert videos.loc["M_0012a.mp4", "part"] == "a"
    assert videos.loc["M_0012_b.avi", "part"] == "b"
    assert videos.loc["M_0013-A.mov", "part"] == "a"  # part letter case does not matter
    assert not videos.loc["random.mp4", "recognized"]


def make_trials(*specs):
    """Minimal subject table for matching tests: (subject_num, sex, n_workbook_rows)."""
    return pd.DataFrame(
        [{"subject_num": n, "subject_id": f"{n:04d}", "sex": s, "n_workbook_rows": k} for n, s, k in specs]
    )


def test_match_videos_statuses(tmp_path):
    touch(tmp_path, "F_0001.mp4", "M_0002.mp4", "F_0003.mp4", "F_0003_copy.mp4", "M_0004b.mp4", "M_0004a.mp4",
          "M_0005a.mp4", "F_0099.mp4", "whatever.mp4")
    videos = scan_videos(tmp_path, PATTERN, EXTENSIONS)
    trials = make_trials((1, "F", 1), (2, "F", 1), (3, "F", 1), (4, "M", 2), (5, "M", 2), (6, "M", 1))

    matched, issues = match_videos(trials, videos)
    status = dict(zip(matched["subject_id"], matched["video_status"]))

    assert status == {
        "0001": "matched",
        "0002": "matched",  # the letter prefix is not a sex code; sex comes from the workbook
        "0003": "matched",  # "F_0003_copy" is not a video name for subject 3 -> unrecognized
        "0004": "matched",  # both parts present
        "0005": "part_mismatch",  # only part a
        "0006": "missing",
    }
    paths = dict(zip(matched["subject_id"], matched["video_paths"]))
    assert [Path(p).name for p in paths["0004"].split(";")] == ["M_0004a.mp4", "M_0004b.mp4"]  # part order
    assert paths["0005"] == ""  # doubtful matches get no path
    assert issues["Videos without a workbook row"] == ["F_0099.mp4: no workbook row for subject 99"]
    assert sorted(issues["Unrecognized video file names"]) == [
        "F_0003_copy.mp4: name does not match the video name pattern",
        "whatever.mp4: name does not match the video name pattern",
    ]


def test_two_files_for_a_single_recording_are_flagged(tmp_path):
    touch(tmp_path, "F_0001.mp4", "F_0001a.mp4")
    matched, issues = match_videos(make_trials((1, "F", 1)), scan_videos(tmp_path, PATTERN, EXTENSIONS))
    assert matched.loc[0, "video_status"] == "duplicate_videos"
    assert len(issues["Video problems"]) == 1


# --- whole step + CLI -------------------------------------------------------------------


@pytest.fixture
def project(tmp_path, monkeypatch):
    """A folder with a synthetic workbook, a videos folder and a .env pointing to both."""
    for name in ("FISH_VIDEO_DIR", "FISH_DB_PATH", "FISH_REFERENCE_PDF", "FISH_OUTPUT_DIR", "FISH_WORKERS", "FISH_CONFIG"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.chdir(tmp_path)
    make_workbook(tmp_path / "db.xlsx", [
        row("0001", "F", "Drug-A", TRACKED),
        row("0002", "M", "Drug-A", TOP_ZERO),
        row("0003", "M", "Drug-B", ALL_NA, uv="10*"),
        row("0003", "M", "Drug-B", ALL_NA, uv="10*"),
        row(None, None, None, ALL_BLANK),  # template row
    ])
    touch(tmp_path / "videos", "F_0001.mp4", "M_0002.mp4", "M_0003a.mp4", "M_0003b.mp4")
    (tmp_path / ".env").write_text("FISH_DB_PATH=db.xlsx\nFISH_VIDEO_DIR=videos\nFISH_OUTPUT_DIR=out\n")
    return tmp_path


def test_build_catalog_writes_all_outputs(project):
    result = build_catalog(load_settings())

    assert not result.has_issues
    assert result.notes["subjects"] == 3 and result.notes["untracked_subjects"] == 1
    for name in ("trials.csv", "videos.csv", "validation_report.md"):
        assert (project / "out" / "catalog" / name).is_file()
    report = (project / "out" / "catalog" / "validation_report.md").read_text()
    assert "Subjects after joining split recordings: 3" in report
    assert "Subject 0003: 2 rows (Excel rows 4;5)" in report


def test_build_catalog_without_video_dir_still_checks_the_workbook(project):
    (project / ".env").write_text("FISH_DB_PATH=db.xlsx\nFISH_OUTPUT_DIR=out\n")

    result = build_catalog(load_settings())

    assert set(result.trials["video_status"]) == {"not_checked"}


def test_cli_validate_and_strict_mode(project, capsys):
    assert main(["validate"]) == 0
    assert "3 subjects" in capsys.readouterr().out

    (project / "videos" / "M_0002.mp4").unlink()  # now one subject has no video
    assert main(["validate"]) == 0  # issues are reported but do not fail by default
    assert main(["validate", "--strict"]) == 1  # ...unless --strict is used


def test_cli_validate_reports_missing_workbook(project, capsys):
    (project / ".env").write_text("FISH_DB_PATH=nope.xlsx\n")
    assert main(["validate"]) == 2
    assert "FISH_DB_PATH points to" in capsys.readouterr().out
